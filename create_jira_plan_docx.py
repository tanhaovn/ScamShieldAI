# -*- coding: utf-8 -*-
"""
Script to generate Jira Sprint Plan docx for SFDS Project.
"""
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_header_footer(doc):
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)
        
        # Header
        hdr = s.header
        p_hdr = hdr.paragraphs[0]
        p_hdr.text = "SFDS - KẾ HOẠCH QUẢN LÝ 6 SPRINT TRÊN JIRA | CAPSTONE PROJECT 1 (CMU-SE 450)"
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for r in p_hdr.runs:
            r.font.name = "Arial"
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(128, 128, 128)
            
        # Footer
        ftr = s.footer
        p_ftr = ftr.paragraphs[0]
        p_ftr.text = "Hệ thống phát hiện ảnh lừa đảo có giải thích (SFDS) - C1SE.06 - Nhóm: Hào, Nhung, Nam, Tuấn"
        p_ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p_ftr.runs:
            r.font.name = "Arial"
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(128, 128, 128)

def format_table(table, col_widths, headers, rows_data):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1F4E79") # Navy Blue
        set_cell_margins(hdr_cells[i], top=140, bottom=140, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

    # Data Rows
    for r_idx, row_values in enumerate(rows_data):
        row = table.add_row()
        cells = row.cells
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_values):
            cells[c_idx].text = str(val)
            set_cell_background(cells[c_idx], bg_color)
            set_cell_margins(cells[c_idx], top=100, bottom=100, left=120, right=120)
            p = cells[c_idx].paragraphs[0]
            
            # Alignments
            if c_idx in [0, 1, 4, 5, 6]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
            for r in p.runs:
                r.font.name = "Arial"
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(30, 41, 59)
                # Specific highlight for status
                if str(val).strip() == "DONE":
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(22, 101, 52) # Dark green
                elif str(val).strip() == "IN PROGRESS":
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(180, 83, 9) # Amber
                elif str(val).strip() == "TO DO":
                    r.font.color.rgb = RGBColor(100, 116, 139) # Gray slate

    # Apply col widths
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = Inches(width)

def build_jira_plan_docx(output_path):
    doc = docx.Document()
    add_header_footer(doc)

    # Document Title Block
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_uni = title_p.add_run("TRƯỜNG ĐẠI HỌC DUY TÂN - VIỆN ĐÀO TẠO QUỐC TẾ (IS)\n")
    r_uni.font.name = "Arial"
    r_uni.font.size = Pt(11)
    r_uni.font.bold = True
    r_uni.font.color.rgb = RGBColor(100, 116, 139)

    r_title = title_p.add_run("KẾ HOẠCH PHÂN CHIA 6 SPRINT TRÊN JIRA\n")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(18)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(31, 78, 121)

    r_sub = title_p.add_run("HỆ THỐNG PHÁT HIỆN VÀ CẢNH BÁO ẢNH LỪA ĐẢO KÈM GIẢI THÍCH (SFDS)\n(Fraudulent Image Detection and Warning System with Explanations)\n")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(12)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(46, 117, 182)

    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(14)
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_meta = meta_p.add_run("Mã môn: CMU-SE 450 | Nhóm: C1SE.06 | Giảng viên hướng dẫn: KS. Hà Mạnh Toàn\nThời gian đồ án: 23/08/2026 - 22/12/2026 | Ngày lập kế hoạch: 15/10/2026")
    r_meta.font.name = "Arial"
    r_meta.font.size = Pt(9.5)
    r_meta.font.italic = True
    r_meta.font.color.rgb = RGBColor(100, 116, 139)

    # Callout Box: Tóm tắt tiến độ hiện tại
    callout_tbl = doc.add_table(rows=1, cols=1)
    callout_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_cell = callout_tbl.rows[0].cells[0]
    set_cell_background(c_cell, "EFF6FF") # Light blue
    set_cell_margins(c_cell, top=140, bottom=140, left=180, right=180)
    c_cell.width = Inches(6.8)
    cp = c_cell.paragraphs[0]
    cp.paragraph_format.space_after = Pt(2)
    c_run1 = cp.add_run("📌 TỔNG QUAN TIẾN ĐỘ THỰC TẾ & ĐỐI CHIẾU VỚI PROPOSAL:\n")
    c_run1.font.name = "Arial"
    c_run1.font.bold = True
    c_run1.font.size = Pt(10)
    c_run1.font.color.rgb = RGBColor(30, 58, 138)
    
    cp2 = c_cell.add_paragraph()
    cp2.paragraph_format.space_after = Pt(2)
    c_run2 = cp2.add_run(
        "• Tài liệu Proposal (v2.0): ĐÃ HOÀN THÀNH (Đã được GVHD & Hội đồng xét duyệt thông qua).\n"
        "• Sprint 1 (02/09 - 16/09/2026): [DONE] Đã hoàn thành 100% (Khởi tạo repo, kiến trúc, DB MySQL schema, thu thập dữ liệu sơ bộ, UI Wireframe).\n"
        "• Sprint 2 (17/09 - 01/10/2026): [DONE] Đã hoàn thành 100% (Tích hợp OCR Tesseract tiếng Việt, mô hình PhishVN URL, Risk Engine tính điểm, UI upload ảnh cơ bản, JWT auth).\n"
        "• Sprint 3 (02/10 - 16/10/2026): [IN PROGRESS / GẦN XONG] Đang hoàn thiện (Visual Forensics EXIF/ELA, Huấn luyện ResNet18, Explanation Builder giải thích lý do, hỗ trợ đa nguồn text/PDF, lưu lịch sử).\n"
        "• Sprint 4 (17/10 - 31/10/2026): [TO DO] Chuẩn bị triển khai (Admin Dashboard hoàn chỉnh, quản lý keyword, OCR Bounding Box highlight trực tiếp trên ảnh, vá bảo mật).\n"
        "• Sprint 5 (01/11 - 15/11/2026): [TO DO] Đánh giá định lượng mô hình (Macro-F1, Recall, CER/WER), tối ưu hiệu năng và bảo mật, CI/CD Docker AWS.\n"
        "• Sprint 6 (17/11 - 01/12/2026): [TO DO] UAT thử nghiệm người dùng, tài liệu kỹ thuật cuối kỳ, hoàn thiện demo và nghiệm thu sản phẩm."
    )
    c_run2.font.name = "Arial"
    c_run2.font.size = Pt(9.5)
    c_run2.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ---------------------------------------------------------
    # PHẦN 1: THIẾT LẬP DỰ ÁN TRÊN JIRA
    # ---------------------------------------------------------
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(4)
    r_h1 = h1.add_run("1. THIẾT LẬP DỰ ÁN TRÊN JIRA (JIRA PROJECT CONFIGURATION)")
    r_h1.font.name = "Arial"
    r_h1.font.bold = True
    r_h1.font.size = Pt(13)
    r_h1.font.color.rgb = RGBColor(31, 78, 121)

    p_cfg = doc.add_paragraph()
    p_cfg.paragraph_format.space_after = Pt(4)
    p_cfg.add_run(
        "Để quản lý đúng chuẩn Scrum theo tài liệu Proposal, khi tạo Project trên Jira bạn cấu hình như sau:\n"
        "• Project Template: Scrum (Software Development)\n"
        "• Project Name: Fraudulent Image Detection and Warning System (SFDS)\n"
        "• Project Key: SFDS\n"
        "• Estimation Statistic: Story Points (1 SP tương đương ~4 giờ công phát triển).\n"
        "• Workflow chuẩn: To Do ➔ In Progress ➔ In Review / Testing ➔ Done\n"
        "• Danh sách Thành viên & Phân vai trên Jira:\n"
        "    1. Phạm Tấn Hào (Leader, Full-stack): Quản trị Backend, Security, Risk Engine, Scrum Master.\n"
        "    2. Đào Thị Mỹ Nhung (UI/UX, QA, Docs): Thiết kế Figma, Tester, Viết tài liệu đồ án.\n"
        "    3. Trần Hoàng Nhật Nam (Frontend, Docs): Lập trình giao diện React/Next.js, tích hợp API, Docs.\n"
        "    4. Nguyễn Phan Đức Tuấn (Full-stack): Lập trình Module AI, Dataset, OCR, API backend."
    )
    for r in p_cfg.runs:
        r.font.name = "Arial"
        r.font.size = Pt(9.5)

    # Bảng Epics
    p_epic_title = doc.add_paragraph()
    p_epic_title.paragraph_format.space_before = Pt(6)
    p_epic_title.paragraph_format.space_after = Pt(4)
    r_ep = p_epic_title.add_run("Danh mục Epics cần tạo trước trên Jira:")
    r_ep.font.name = "Arial"
    r_ep.font.bold = True
    r_ep.font.size = Pt(10.5)
    r_ep.font.color.rgb = RGBColor(46, 117, 182)

    epic_table = doc.add_table(rows=1, cols=4)
    epic_widths = [1.2, 1.8, 2.8, 1.0]
    epic_headers = ["Epic Key", "Tên Epic (Epic Name)", "Mục tiêu & Phạm vi chức năng", "Liên kết FR/NFR"]
    epic_rows = [
        ["EPIC-1", "Hạ tầng & Kiến trúc Hệ thống", "Khởi tạo repo, base server FastAPI, MySQL DB schema, Docker & CI/CD", "NFR01, NFR03"],
        ["EPIC-2", "Lõi Xử lý Ảnh & OCR Tiếng Việt", "Tích hợp Tesseract OCR tiếng Việt, bóc tách text, bounding box, quét QR & PDF", "FR01, FR02"],
        ["EPIC-3", "Mô hình AI & Đánh giá Rủi ro", "Tích hợp PhishVN (URL), visual forensics (EXIF/ELA), model ResNet18, Risk Engine", "FR02, FR03"],
        ["EPIC-4", "Giải thích & Khuyến nghị An toàn", "Bóc tách lý do cảnh báo, highlight bằng chứng, khuyến nghị hành động an toàn", "FR03, FR04"],
        ["EPIC-5", "Giao diện Web Người dùng & Admin", "UI kiểm tra ảnh cho khách, lưu history localStorage, Dashboard quản trị DB", "FR01, FR05, FR06"],
        ["EPIC-6", "Kiểm thử, Đánh giá & Hoàn thiện", "Unit test, UAT, đánh giá Macro-F1, Recall, CER/WER, tài liệu kỹ thuật cuối kỳ", "NFR01, NFR02"]
    ]
    format_table(epic_table, epic_widths, epic_headers, epic_rows)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ---------------------------------------------------------
    # PHẦN 2: BẢNG TỔNG QUAN 6 SPRINTS THEO PROPOSAL
    # ---------------------------------------------------------
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(4)
    r_h2 = h2.add_run("2. BẢNG TỔNG QUAN LỊCH TRÌNH 6 SPRINT (THEO PROPOSAL TABLE 8)")
    r_h2.font.name = "Arial"
    r_h2.font.bold = True
    r_h2.font.size = Pt(13)
    r_h2.font.color.rgb = RGBColor(31, 78, 121)

    sprint_summary_table = doc.add_table(rows=1, cols=7)
    sprint_sum_widths = [0.8, 1.2, 1.0, 1.0, 0.8, 1.0, 1.0]
    sprint_sum_headers = ["Sprint", "Tên Sprint trên Jira", "Bắt đầu", "Kết thúc", "Giờ công", "Tổng SP", "Trạng thái"]
    sprint_sum_rows = [
        ["Sprint 1", "Khởi tạo nền tảng & Wireframe", "02/09/2026", "16/09/2026", "200 hrs", "50 SP", "DONE"],
        ["Sprint 2", "Lõi OCR & PhishVN URL Scanner", "17/09/2026", "01/10/2026", "200 hrs", "50 SP", "DONE"],
        ["Sprint 3", "Visual Forensics & Explanation", "02/10/2026", "16/10/2026", "200 hrs", "50 SP", "IN PROGRESS"],
        ["Sprint 4", "Admin Dashboard & OCR Highlight", "17/10/2026", "31/10/2026", "200 hrs", "50 SP", "TO DO"],
        ["Sprint 5", "Đánh giá AI & Tối ưu Bảo mật", "01/11/2026", "15/11/2026", "200 hrs", "50 SP", "TO DO"],
        ["Sprint 6", "UAT, Đóng gói Demo & Báo cáo", "17/11/2026", "01/12/2026", "200 hrs", "50 SP", "TO DO"]
    ]
    format_table(sprint_summary_table, sprint_sum_widths, sprint_sum_headers, sprint_sum_rows)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ---------------------------------------------------------
    # PHẦN 3: CHI TIẾT TỪNG SPRINT ĐỂ TẠO TRỰC TIẾP TRÊN JIRA
    # ---------------------------------------------------------
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(12)
    h3.paragraph_format.space_after = Pt(4)
    r_h3 = h3.add_run("3. CHI TIẾT CÁC USER STORIES & TASKS CHO TỪNG SPRINT")
    r_h3.font.name = "Arial"
    r_h3.font.bold = True
    r_h3.font.size = Pt(13)
    r_h3.font.color.rgb = RGBColor(31, 78, 121)

    sprints_detail = [
        {
            "id": "Sprint 1",
            "name": "SFDS Sprint 1: Khởi tạo Nền tảng, DB Schema & Wireframe UI",
            "goal": "Thiết lập toàn bộ khung dự án (backend, frontend, DB), thu thập dataset ban đầu và thiết kế hoàn chỉnh wireframe UI/UX.",
            "dates": "02/09/2026 - 16/09/2026 (13 ngày, 200 giờ công, 50 SP)",
            "status": "DONE (Đã hoàn thành 100%)",
            "tasks": [
                ["SFDS-1", "Story", "EPIC-1", "Khởi tạo cấu trúc Backend FastAPI & Cấu hình môi trường", "Setup project FastAPI, file .env.example, requirements.txt, cấu trúc thư mục module hóa.", "Phạm Tấn Hào", 8, "DONE"],
                ["SFDS-2", "Task", "EPIC-1", "Thiết kế & Tạo Database Schema MySQL (schema.sql)", "Viết schema.sql chứa bảng admin, categories, scans, datasets, evaluations, audit_logs.", "Phạm Tấn Hào", 8, "DONE"],
                ["SFDS-3", "Story", "EPIC-5", "Thiết kế Wireframe & Prototype Figma UI/UX", "Thiết kế giao diện người dùng kiểm tra ảnh và Dashboard quản trị admin trên Figma.", "Đào Thị Mỹ Nhung", 8, "DONE"],
                ["SFDS-4", "Task", "EPIC-3", "Khảo sát và thu thập bộ Dataset ảnh & URL lừa đảo ban đầu", "Thu thập mẫu ảnh hoá đơn giả, tin nhắn SMS lừa đảo, tải dataset PhishVN từ Mendeley.", "Nguyễn Phan Đức Tuấn", 8, "DONE"],
                ["SFDS-5", "Task", "EPIC-1", "Khởi tạo Frontend Project (React/Vite & HTML/JS cơ bản)", "Tạo cấu trúc frontend, tích hợp component upload ảnh và preview, script khởi động BAT.", "Trần Hoàng Nhật Nam", 8, "DONE"],
                ["SFDS-6", "Task", "EPIC-6", "Thiết lập Repository GitHub & Quy chuẩn Git Flow", "Tạo repo GitHub, quy định commit message, branch dev/main, cấu hình .gitignore.", "Trần Hoàng Nhật Nam", 5, "DONE"],
                ["SFDS-7", "Task", "EPIC-6", "Xây dựng kịch bản kiểm thử nền tảng và tài liệu khởi động", "Viết hướng dẫn cài đặt trong README.md, script SETUP_PROJECT.bat và START_PROJECT.bat.", "Đào Thị Mỹ Nhung", 5, "DONE"]
            ]
        },
        {
            "id": "Sprint 2",
            "name": "SFDS Sprint 2: Tích hợp Lõi OCR Tiếng Việt, PhishVN & Risk Engine Sơ bộ",
            "goal": "Người dùng có thể upload ảnh, trích xuất text tiếng Việt bằng Tesseract, phát hiện URL độc hại qua PhishVN và nhận điểm đánh giá sơ bộ.",
            "dates": "17/09/2026 - 01/10/2026 (13 ngày, 200 giờ công, 50 SP)",
            "status": "DONE (Đã hoàn thành 100%)",
            "tasks": [
                ["SFDS-8", "Story", "EPIC-2", "Xây dựng Module OCR Tesseract tiếng Việt (tessdata/vie)", "Tích hợp pytesseract, đọc ảnh JPG/PNG, xử lý chữ tiếng Việt có dấu và số tài khoản ngân hàng.", "Nguyễn Phan Đức Tuấn", 8, "DONE"],
                ["SFDS-9", "Story", "EPIC-3", "Tích hợp mô hình PhishVN v3.1.0 phân loại URL độc hại", "Tách URL từ OCR, trích xuất 21 đặc trưng lexical, inference mô hình HistGradientBoosting.", "Nguyễn Phan Đức Tuấn", 8, "DONE"],
                ["SFDS-10", "Story", "EPIC-3", "Xây dựng Rule-based Risk Engine tính điểm rủi ro cơ sở", "So khớp từ khóa nhạy cảm tiếng Việt (ngân hàng, OTP, trúng thưởng), tính điểm 0-100.", "Phạm Tấn Hào", 8, "DONE"],
                ["SFDS-11", "Story", "EPIC-5", "Xây dựng Giao diện Upload ảnh & Hiển thị kết quả sơ bộ", "User upload kéo thả ảnh, chọn category, hiển thị thanh điểm rủi ro và 3 nhãn cảnh báo.", "Trần Hoàng Nhật Nam", 8, "DONE"],
                ["SFDS-12", "Story", "EPIC-1", "Xây dựng Cơ chế Đăng nhập Admin bằng JWT & Quản trị phiên", "Endpoint /token và /me, mã hóa mật khẩu PBKDF2-HMAC-SHA256, bảo vệ các API admin.", "Phạm Tấn Hào", 8, "DONE"],
                ["SFDS-13", "Task", "EPIC-6", "Viết Unit Test cho OCR và API /check-image sơ bộ", "Kiểm thử API /check-image với mock dữ liệu, test trường hợp không có MySQL vẫn chạy an toàn.", "Đào Thị Mỹ Nhung", 5, "DONE"],
                ["SFDS-14", "Task", "EPIC-5", "Hoàn thiện Swagger API Documentation (/docs)", "Cấu hình OpenAPI schema, mô tả đầy đủ tham số request/response cho các endpoint.", "Trần Hoàng Nhật Nam", 5, "DONE"]
            ]
        },
        {
            "id": "Sprint 3",
            "name": "SFDS Sprint 3: Visual Forensics, Huấn luyện ResNet18 & Explanation Builder",
            "goal": "Bổ sung phân tích ảnh chuyên sâu (EXIF/ELA), huấn luyện mô hình phân loại ảnh ResNet18, bóc tách giải thích nguyên nhân và khuyến nghị an toàn.",
            "dates": "02/10/2026 - 16/10/2026 (13 ngày, 200 giờ công, 50 SP)",
            "status": "IN PROGRESS (Đang hoàn thiện / Gần xong)",
            "tasks": [
                ["SFDS-15", "Story", "EPIC-3", "Xây dựng Module Visual Forensics (EXIF Metadata & ELA)", "Phân tích chỉnh sửa ảnh bằng Error Level Analysis (ELA), đọc metadata camera/phần mềm sửa.", "Nguyễn Phan Đức Tuấn", 8, "IN PROGRESS"],
                ["SFDS-16", "Story", "EPIC-3", "Huấn luyện & Tích hợp Mô hình Phân loại Ảnh ResNet18", "Xây dựng train_image_model.py, fine-tune ResNet18 trên bộ ảnh phishing, xuất file pt.", "Nguyễn Phan Đức Tuấn", 8, "IN PROGRESS"],
                ["SFDS-17", "Story", "EPIC-4", "Phát triển Explanation Builder & Khuyến nghị Hành động", "Bóc tách nguyên nhân rủi ro thành 4 nhóm tín hiệu, tự động sinh lời khuyên an toàn theo ngữ cảnh.", "Phạm Tấn Hào", 8, "DONE"],
                ["SFDS-18", "Story", "EPIC-2", "Bổ sung Hỗ trợ Quét QR Code và Trích xuất Tài liệu PDF", "Giải mã mã QR chứa link/thanh toán trong ảnh, hỗ trợ phân tích file PDF tối đa 10 trang.", "Phạm Tấn Hào", 8, "DONE"],
                ["SFDS-19", "Story", "EPIC-5", "Xây dựng Lịch sử Quét bằng localStorage cho Khách", "Lưu tối đa 50 lượt quét gần nhất tại trình duyệt, xem lại chi tiết không cần tạo tài khoản.", "Trần Hoàng Nhật Nam", 8, "DONE"],
                ["SFDS-20", "Task", "EPIC-5", "Cập nhật Giao diện hiển thị Bóc tách Điểm & Khuyến nghị", "Hiển thị Breakdown rủi ro (Nội dung OCR, URL, Ảnh, Model) và hộp cảnh báo an toàn trực quan.", "Trần Hoàng Nhật Nam", 5, "DONE"],
                ["SFDS-21", "Task", "EPIC-6", "Kiểm thử Tích hợp End-to-End quy trình Quét đa nguồn", "Test upload ảnh, paste text, URL trực tiếp và file PDF; ghi nhận bug phát sinh.", "Đào Thị Mỹ Nhung", 5, "IN PROGRESS"]
            ]
        },
        {
            "id": "Sprint 4",
            "name": "SFDS Sprint 4: Hoàn thiện Admin Dashboard, OCR Bounding Box & Security",
            "goal": "Hoàn thiện toàn bộ các trang chức năng Admin (quản lý dataset, keyword, audit log), nâng cấp OCR với bounding box highlight và vá lỗ hổng bảo mật.",
            "dates": "17/10/2026 - 31/10/2026 (13 ngày, 200 giờ công, 50 SP)",
            "status": "TO DO (Chưa thực hiện)",
            "tasks": [
                ["SFDS-22", "Story", "EPIC-5", "Phát triển Trang Quản trị Admin: Quản lý Dataset & AI Models", "Giao diện quản lý danh sách tập dữ liệu huấn luyện, thông tin model checkpoints và logs.", "Trần Hoàng Nhật Nam", 8, "TO DO"],
                ["SFDS-23", "Story", "EPIC-5", "Quản lý Danh mục Từ khóa Scam động trên Admin UI", "Cho phép Admin thêm/sửa/xóa bộ từ khóa lừa đảo và trọng số rủi ro thay vì hardcode.", "Trần Hoàng Nhật Nam", 8, "TO DO"],
                ["SFDS-24", "Story", "EPIC-2", "Nâng cấp OCR: Highlight Bounding Box vùng nghi ngờ trên ảnh", "Tọa độ trích xuất từ Tesseract được gửi về frontend để vẽ khung đỏ đánh dấu từ ngữ gian lận.", "Nguyễn Phan Đức Tuấn", 8, "TO DO"],
                ["SFDS-25", "Story", "EPIC-1", "Vá Lỗ hổng Bảo mật API & Kiểm tra Magic Bytes File Upload", "Khắc phục các điểm trong REVIEW.md: Validate file bytes thực tế, chặn file độc, cấu hình JWT an toàn.", "Phạm Tấn Hào", 8, "TO DO"],
                ["SFDS-26", "Task", "EPIC-1", "Tối ưu Lưu trữ Lịch sử Quét & Audit Log cho Admin trong MySQL", "Lưu vết truy cập của Admin, xuất báo cáo lịch sử quét ra định dạng CSV/Excel.", "Phạm Tấn Hào", 8, "TO DO"],
                ["SFDS-27", "Task", "EPIC-6", "Kiểm thử Bảo mật (Security & Authorization Testing)", "Test kiểm tra quyền truy cập trái phép API admin, test upload file giả mạo đuôi ảnh.", "Đào Thị Mỹ Nhung", 5, "TO DO"],
                ["SFDS-28", "Task", "EPIC-6", "Cải tiến UX: Loading Skeleton, Toast Notification thay thế Alert", "Thay thế hoàn toàn window.alert bằng toast messages, loading state mượt mà khi AI phân tích.", "Đào Thị Mỹ Nhung", 5, "TO DO"]
            ]
        },
        {
            "id": "Sprint 5",
            "name": "SFDS Sprint 5: Đánh giá Định lượng AI, Tối ưu Hiệu năng & CI/CD Docker",
            "goal": "Đo đạc các chỉ số khoa học cam kết trong Proposal (Macro-F1, Recall, False Alarm, CER/WER), tối ưu tốc độ phản hồi và đóng gói Docker CI/CD.",
            "dates": "01/11/2026 - 15/11/2026 (13 ngày, 200 giờ công, 50 SP)",
            "status": "TO DO (Chưa thực hiện)",
            "tasks": [
                ["SFDS-29", "Story", "EPIC-3", "Đánh giá Khoa học Mô hình PhishVN & ResNet18 (Macro-F1, Recall)", "Thực hiện kiểm thử trên test set độc lập, xuất Confusion Matrix, tính Macro-F1, Recall.", "Nguyễn Phan Đức Tuấn", 8, "TO DO"],
                ["SFDS-30", "Story", "EPIC-2", "Benchmark Độ chính xác OCR Tiếng Việt (Đo chỉ số CER & WER)", "Tạo tập test ground-truth tiếng Việt, đo Character Error Rate (CER) và Word Error Rate (WER).", "Nguyễn Phan Đức Tuấn", 8, "TO DO"],
                ["SFDS-31", "Story", "EPIC-3", "Hiệu chuẩn Ngưỡng Điểm (Threshold Calibration) Giảm False Positive", "Tinh chỉnh ngưỡng an toàn (25) và nguy hiểm (60) trên tập dữ liệu ảnh hoá đơn/tin nhắn hợp lệ.", "Phạm Tấn Hào", 8, "TO DO"],
                ["SFDS-32", "Story", "EPIC-1", "Đóng gói Toàn bộ Hệ thống bằng Docker & Docker Compose", "Viết Dockerfile cho Backend, Frontend, cấu hình docker-compose chạy cùng MySQL & Redis.", "Phạm Tấn Hào", 8, "TO DO"],
                ["SFDS-33", "Task", "EPIC-1", "Thiết lập Pipeline CI/CD Tự động trên GitHub Actions", "Tự động chạy pytest kiểm thử và kiểm tra build frontend mỗi khi có commit mới.", "Trần Hoàng Nhật Nam", 5, "TO DO"],
                ["SFDS-34", "Task", "EPIC-6", "Stress Test & Đo Đạc Hiệu năng Tốc độ Phản hồi Hệ thống", "Dùng Postman/Locust kiểm thử tải nhiều người dùng upload cùng lúc, tối ưu thời gian xử lý < 3s.", "Đào Thị Mỹ Nhung", 8, "TO DO"],
                ["SFDS-35", "Task", "EPIC-6", "Đánh giá Mức độ Hữu ích của Lời giải thích (Explanation Usefulness)", "Khảo sát và ghi nhận tỉ lệ lý do cảnh báo được người dùng chấp nhận (thang điểm 1-5).", "Đào Thị Mỹ Nhung", 5, "TO DO"]
            ]
        },
        {
            "id": "Sprint 6",
            "name": "SFDS Sprint 6: Kiểm thử UAT, Triển khai Demo & Hoàn thiện Tài liệu Cuối kỳ",
            "goal": "Nghiệm thu toàn diện cùng người dùng (UAT), fix bug tồn đọng, triển khai demo và hoàn tất toàn bộ bộ tài liệu bàn giao đồ án Capstone 1.",
            "dates": "17/11/2026 - 01/12/2026 (13 ngày, 200 giờ công, 50 SP)",
            "status": "TO DO (Chưa thực hiện)",
            "tasks": [
                ["SFDS-36", "Story", "EPIC-6", "Tổ chức Kiểm thử Chấp nhận Người dùng (UAT - User Acceptance Testing)", "Mời người dùng thử nghiệm upload các kịch bản lừa đảo thực tế, tổng hợp phản hồi và phiếu lỗi.", "Đào Thị Mỹ Nhung", 8, "TO DO"],
                ["SFDS-37", "Story", "EPIC-1", "Triển khai Môi trường Demo Thử nghiệm (AWS EC2 / Demo Server)", "Deploy hệ thống hoàn chỉnh lên server demo phục vụ báo cáo với Giảng viên và Hội đồng.", "Phạm Tấn Hào", 8, "TO DO"],
                ["SFDS-38", "Bug", "EPIC-6", "Bug Fixing Tổng thể & Đánh bóng Giao diện Đồ án", "Khắc phục triệt để các lỗi giao diện, lỗi tràn chữ, lỗi bộ nhớ khi xử lý file lớn.", "Trần Hoàng Nhật Nam", 8, "TO DO"],
                ["SFDS-39", "Task", "EPIC-6", "Biên soạn Tài liệu Hướng dẫn Sử dụng (User Manual & Admin Manual)", "Soạn tài liệu hướng dẫn trực quan kèm ảnh chụp màn hình từng bước thao tác cho người dùng.", "Đào Thị Mỹ Nhung", 8, "TO DO"],
                ["SFDS-40", "Task", "EPIC-6", "Hoàn thiện Báo cáo Kỹ thuật Đồ án (Technical Report Capstone 1)", "Viết hoàn chỉnh tài liệu kiến trúc, sơ đồ thực thể ERD, thiết kế chi tiết API và kết quả thực nghiệm.", "Trần Hoàng Nhật Nam", 8, "TO DO"],
                ["SFDS-41", "Task", "EPIC-6", "Chuẩn bị Kịch bản Báo cáo & Video Demo Đồ án", "Xây dựng slide thuyết trình, kịch bản demo trực tiếp các case study nhận diện lừa đảo thành công.", "Phạm Tấn Hào", 5, "TO DO"],
                ["SFDS-42", "Task", "EPIC-6", "Nghiệm thu Codebase, Đóng gói Release Version v1.0", "Gắn tag Git release v1.0, sao lưu database và kiểm tra toàn bộ checklist bàn giao.", "Nguyễn Phan Đức Tuấn", 5, "TO DO"]
            ]
        }
    ]

    col_widths = [0.8, 0.7, 0.8, 2.5, 0.9, 0.5, 0.6]
    task_headers = ["Issue Key", "Loại", "Epic", "Tóm tắt Công việc (Summary)", "Assignee", "SP", "Status"]

    for sp in sprints_detail:
        sp_p = doc.add_paragraph()
        sp_p.paragraph_format.space_before = Pt(14)
        sp_p.paragraph_format.space_after = Pt(2)
        r_sp = sp_p.add_run(f"3.{sp['id'][-1]}. {sp['name']}")
        r_sp.font.name = "Arial"
        r_sp.font.bold = True
        r_sp.font.size = Pt(11.5)
        r_sp.font.color.rgb = RGBColor(31, 78, 121)

        sp_info = doc.add_paragraph()
        sp_info.paragraph_format.space_after = Pt(4)
        sp_info.add_run(f"• Thời gian: {sp['dates']} | Trạng thái hiện tại: ").font.bold = True
        r_st = sp_info.add_run(f"{sp['status']}\n")
        r_st.font.bold = True
        if "DONE" in sp['status']:
            r_st.font.color.rgb = RGBColor(22, 101, 52)
        elif "IN PROGRESS" in sp['status']:
            r_st.font.color.rgb = RGBColor(180, 83, 9)
        else:
            r_st.font.color.rgb = RGBColor(100, 116, 139)
            
        sp_info.add_run(f"• Mục tiêu Sprint (Sprint Goal): {sp['goal']}")
        for r in sp_info.runs:
            r.font.name = "Arial"
            r.font.size = Pt(9.5)

        table_data = []
        for t in sp['tasks']:
            # t: [key, type, epic, summary, desc, assignee, sp, status]
            table_data.append([t[0], t[1], t[2], f"{t[3]}\n(Mô tả: {t[4]})", t[5].split()[-1], t[6], t[7]])

        tbl = doc.add_table(rows=1, cols=7)
        format_table(tbl, col_widths, task_headers, table_data)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ---------------------------------------------------------
    # PHẦN 4: HƯỚNG DẪN THAO TÁC TẠO TRỰC TIẾP TRÊN JIRA
    # ---------------------------------------------------------
    h4 = doc.add_paragraph()
    h4.paragraph_format.space_before = Pt(14)
    h4.paragraph_format.space_after = Pt(4)
    r_h4 = h4.add_run("4. HƯỚNG DẪN THAO TÁC NHẬP DỮ LIỆU LÊN JIRA TỪ TÀI LIỆU NÀY")
    r_h4.font.name = "Arial"
    r_h4.font.bold = True
    r_h4.font.size = Pt(13)
    r_h4.font.color.rgb = RGBColor(31, 78, 121)

    steps_p = doc.add_paragraph()
    steps_p.paragraph_format.space_after = Pt(4)
    steps_p.add_run(
        "Bạn chỉ cần mở trang Jira của nhóm và thực hiện lần lượt theo 5 bước dưới đây:\n\n"
        "Bước 1: Tạo Epics trong mục Backlog hoặc Timeline\n"
        "  - Vào Jira > chọn Project SFDS > mở menu Backlog (hoặc Timeline).\n"
        "  - Chọn nút '+ Create Epic' và tạo lần lượt 6 Epics từ EPIC-1 đến EPIC-6 theo bảng ở Mục 1.\n\n"
        "Bước 2: Tạo sẵn 6 Sprints trong Backlog\n"
        "  - Trong màn hình Backlog, nhấp nút 'Create Sprint' 6 lần để tạo Sprint 1 đến Sprint 6.\n"
        "  - Nhấp vào dấu ba chấm (...) ở góc mỗi Sprint > chọn 'Edit Sprint':\n"
        "      + Nhập Sprint Name, Start Date, End Date và Sprint Goal tương ứng theo từng bảng ở Mục 3.\n\n"
        "Bước 3: Tạo các User Stories, Tasks và Gán vào Sprint tương ứng\n"
        "  - Nhấp '+ Create Issue' dưới mỗi Sprint (hoặc nút Create to ở thanh menu trên cùng):\n"
        "      + Issue Type: Chọn Story / Task / Bug theo cột 'Loại'.\n"
        "      + Summary: Copy tiêu đề công việc.\n"
        "      + Description: Dán mô tả chi tiết và tiêu chí hoàn thành.\n"
        "      + Epic Link: Chọn đúng Epic tương ứng (EPIC-1 đến EPIC-6).\n"
        "      + Assignee: Gán cho thành viên nhóm (Hào, Nhung, Nam, Tuấn).\n"
        "      + Story Points: Điền số điểm (5 hoặc 8 điểm).\n\n"
        "Bước 4: Cập nhật Trạng thái (Status) theo Hiện trạng Thực tế\n"
        "  - Đối với Sprint 1: Bấm 'Start Sprint' > Kéo toàn bộ các task sang cột 'DONE' > Bấm 'Complete Sprint'.\n"
        "  - Đối với Sprint 2: Bấm 'Start Sprint' > Kéo toàn bộ các task sang cột 'DONE' > Bấm 'Complete Sprint'.\n"
        "  - Đối với Sprint 3: Bấm 'Start Sprint' (đang là Active Sprint hiện tại). Kéo các task SFDS-17, 18, 19, 20 sang 'DONE', các task SFDS-15, 16, 21 để ở 'IN PROGRESS'.\n"
        "  - Đối với Sprint 4, Sprint 5, Sprint 6: Giữ nguyên trong Backlog ở trạng thái 'TO DO'.\n\n"
        "Bước 5: Theo dõi biểu đồ Burndown Chart & Báo cáo với GVHD\n"
        "  - Jira sẽ tự động sinh Burndown chart, Velocity chart và Cumulative Flow Diagram rất đẹp mắt, sẵn sàng chụp ảnh đưa vào Slide bảo vệ và Báo cáo tiến độ!"
    )
    for r in steps_p.runs:
        r.font.name = "Arial"
        r.font.size = Pt(9.5)

    doc.save(output_path)
    print(f"Successfully generated Jira plan Word docx at: {output_path}")

if __name__ == "__main__":
    output_file = r"d:\scam-detector\KE_HOACH_CHIA_SPRINT_JIRA_SFDS.docx"
    build_jira_plan_docx(output_file)
